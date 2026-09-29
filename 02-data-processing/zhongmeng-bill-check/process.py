# -*- coding: utf-8 -*-
"""中盟 SOP 的 Excel 加工、报价和天图任务生成。

输入账单按 SOP 使用 A/F/I/L/O/Q 列；报价表支持“渠道 sheet + 分区列”或
“费用名称/价格”二维表。所有写回均为值，避免跨工作簿链接失效。
"""
from __future__ import annotations
import csv, re, shutil, unicodedata
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

LETTERS = {c: i for i, c in enumerate("ABCDEFGHIJKLMNOPQRSTUVWXYZ", 1)}
def col(c): return LETTERS.get(str(c).upper(), int(c) if str(c).isdigit() else 1)
def col_letter(n):
    out=""
    while n:
        n, rem = divmod(int(n)-1, 26); out=chr(65+rem)+out
    return out
def norm(v): return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", str(v or "").strip().lower())
def num(v, default=None):
    if v is None or v == "": return default
    try: return float(str(v).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        m = re.search(r"-?\d+(?:\.\d+)?", str(v)); return float(m.group()) if m else default
def waybill(v):
    s = str(v or "").strip()
    return re.split(r"[_-]", s, maxsplit=1)[0].strip()

class QuoteBook:
    def __init__(self, path):
        # 报价表很小但会被数据列表的数万行反复查询；普通模式支持常数时间随机取格。
        self.path = Path(path); self.wb = load_workbook(self.path, data_only=True, read_only=False)
        self.sheets = {norm(s): s for s in self.wb.sheetnames}
        self._lookup_cache = {}
    def sheet(self, channel):
        n = norm(channel)
        if n in self.sheets: return self.sheets[n]
        for k, s in self.sheets.items():
            if n and (n in k or k in n): return s
        # 不允许静默退到目录页；HWT 等报价表里没有的渠道必须明确记为未匹配。
        return None
    @staticmethod
    def _zone_col(zone):
        z=int(num(zone, 0) or 0)
        return {2:5,3:6,4:6,5:7,6:7,7:8,8:8}.get(z)
    @staticmethod
    def _base_zone_col(zone):
        z=int(num(zone, 0) or 0)
        return z+2 if 2 <= z <= 8 else None
    def lookup(self, channel, fee, zone=None, weight=None):
        key=(norm(channel),norm(fee),num(zone),num(weight))
        if key in self._lookup_cache: return self._lookup_cache[key]
        result=self._lookup_uncached(channel, fee, zone, weight)
        self._lookup_cache[key]=result
        return result
    def _lookup_uncached(self, channel, fee, zone=None, weight=None):
        target=norm(fee); sn=self.sheet(channel)
        # 报价文件没有 HWT 独立 sheet；仅对明确的 HWT 渠道按现有报价规则复用商业附加费表。
        if not sn and "hwt" in norm(channel) and target != norm("运费"):
            sn=self.sheet("UPS@Gr-GLA商业")
        ws = self.wb[sn] if sn else None
        if ws is None: return None
        # 真实报价表固定版式：附加费标题在 B/C，分区价在 E:H。
        fee_rows={"住宅旺季附加费":158,"住宅地址":158,"超尺寸附加费":159,"超尺寸体积":159,"超重附加费":160,
                  "不规则包装":162,"超大件":163,"补收住宅超大件":163,"商业超大件":164}
        matched=None
        for key,rowno in fee_rows.items():
            if norm(key) in target or target in norm(key): matched=rowno; break
        if matched:
            if matched==158: return round(num(ws.cell(matched,5).value),2)
            zc=self._zone_col(zone)
            return round(num(ws.cell(matched,zc).value),2) if zc and num(ws.cell(matched,zc).value) is not None else None
        if norm("2类偏远") in target: return round(num(ws.cell(168,5).value),2)
        # SOP 运费1/2：用计费重精确匹配报价表 C 列，按分区 2-8 选择 D:J。
        w = num(weight)
        zc=self._base_zone_col(zone)
        if w is not None and zc:
            lookup=round(w,2)
            for r in range(4, min(ws.max_row,153)+1):
                rw=num(ws.cell(r,3).value)
                if rw is not None and round(rw,2)==lookup:
                    price=num(ws.cell(r,zc).value)
                    return round(price,2) if price is not None else None
        return None

RULES = [
 ("1类偏远", "偏远", "fixed", "remote_1"), ("超尺寸附加费", "超长", "quote", None),
 ("地址修正费", "address", "per_piece_fixed", "address_correction"), ("超大件", "超长", "quote", None),
 ("2类偏远", "超偏远", "quote", None), ("住宅旺季附加费", ["住宅私人"], "quote_piece", None),
 ("不规则包装", None, "quote", None), ("住宅地址", "住宅私人", "quote_piece", None),
 ("运费", None, "freight", None), ("补收住宅地址", "住宅私人", "per_piece_fixed", "residential_rebill"),
 ("补收住宅超大件", ["住宅私人", "超长"], "quote", None), ("重复使用面单-附加费", None, "per_piece_fixed", "reuse_label"),
 ("拍照", None, "per_piece_fixed", "photo"), ("超尺寸-体积", "超长", "quote", None),
 ("商业超大件", "超长", "quote", None), ("POD", "pod", "per_piece_fixed", "pod"),
 ("超重附加费", "超重", "quote", None),
]
def match_rule(fee):
    n = norm(fee)
    for name, mark, kind, key in RULES:
        if norm(name) in n or n in norm(name): return name, mark, kind, key
    return None

def backup(path):
    p = Path(path); base = p.with_name(p.stem + "_原始备份.xlsx"); out = base; i = 2
    while out.exists(): out = p.with_name(f"{p.stem}_原始备份_{i}.xlsx"); i += 1
    shutil.copy2(p, out); return out

def _header_col(ws, names, fallback):
    wanted={norm(x) for x in names}
    for c in range(1, min(ws.max_column, 40)+1):
        if norm(ws.cell(1,c).value) in wanted:
            return c
    return fallback

def process_data_list(path, quote_path, dry_run=False):
    """按 SOP 运费2处理数据列表并写数值，避免交付文件依赖外部公式链接。"""
    if not path or not Path(path).is_file(): return 0
    wb=load_workbook(path, data_only=False)
    qb=QuoteBook(quote_path); changed=0
    for ws in wb.worksheets:
        fee_col=_header_col(ws,("费用名","费用名称"),None)
        channel_col=_header_col(ws,("渠道","渠道名称"),1)
        zone_col=_header_col(ws,("分区","区域"),5)
        weight_col=_header_col(ws,("计费重","单件计费重"),8)
        amount_col=_header_col(ws,("金额","费用金额"),11)
        for r in range(2, ws.max_row+1):
            if fee_col and norm(ws.cell(r,fee_col).value) != norm("运费"): continue
            channel,zone,weight=ws.cell(r,channel_col).value,ws.cell(r,zone_col).value,ws.cell(r,weight_col).value
            # 清掉旧版本留下的外部链接公式，即使本行这次因数据缺失无法匹配。
            if not dry_run:
                ws.cell(r,21).value=None
                ws.cell(r,22).value=None
            if channel in (None,""): continue
            rate=qb.lookup(channel,"运费",zone,weight)
            if rate is None: continue
            original=num(ws.cell(r,amount_col).value)
            if not dry_run:
                ws.cell(r,21).value=round(rate,2)
                if original is not None:
                    ws.cell(r,22).value=round(rate-original,2)
            changed+=1
    if changed and not dry_run:
        backup(path)
        wb.save(path)
    return changed

def process(path, quote_path, cfg, log=print, dry_run=False, zone_index=None):
    path, quote_path = Path(path), Path(quote_path)
    if not dry_run and cfg.get("run", {}).get("backup_original", True): bkp = backup(path)
    else: bkp = None
    wb = load_workbook(path); ws = wb.active; qb = QuoteBook(quote_path)
    C = cfg.get("columns", {}); ci = {k: col(v) for k,v in C.items()}
    for key,title in (("zone","分区"),("quote","报价表"),("diff","差异")): ws.cell(1,ci[key]).value=title
    # Q 列保存纯数字分区；导出或转发账单时不依赖数据列表文件路径。
    for r in range(2, ws.max_row + 1):
        no=waybill(ws.cell(r,ci["waybill"]).value)
        existing_zone=ws.cell(r,ci["zone"]).value
        if isinstance(existing_zone,str) and existing_zone.startswith("="):
            ws.cell(r,ci["zone"]).value=None
        existing_diff=ws.cell(r,ci["diff"]).value
        if isinstance(existing_diff,str) and existing_diff.startswith("="):
            ws.cell(r,ci["diff"]).value=None
        if not no: continue
        zone=zone_index.get(no) if zone_index else None
        if zone is not None: ws.cell(r,ci["zone"]).value=int(zone)
    fees=[norm(ws.cell(r,ci["fee"]).value) for r in range(2,ws.max_row+1)]
    skip_freight="运费更正手续费" in fees
    stats = {"总行数": 0, "已报价": 0, "未匹配报价": 0, "缺少分区": 0, "跳过运费": 0, "差异超容差": 0, "天图任务": 0, "未匹配原因": {}, "强制登记": []}
    tasks = []
    for r in range(2, ws.max_row + 1):
        stats["总行数"] += 1; fee = ws.cell(r, ci["fee"]).value; amount = num(ws.cell(r, ci["amount"]).value, 0)
        if not fee or (cfg.get("run", {}).get("skip_zero_amount", True) and not amount): continue
        rule = match_rule(fee)
        if not rule: continue
        name, mark, kind, key = rule; channel = ws.cell(r, ci["channel"]).value; no=waybill(ws.cell(r,ci["waybill"]).value)
        zone=ws.cell(r,ci["zone"]).value
        if (zone is None or zone == "" or (isinstance(zone, str) and zone.startswith("="))) and zone_index:
            zone=zone_index.get(no)
            if zone is not None: ws.cell(r,ci["zone"]).value=int(zone)
        pieces = num(ws.cell(r, ci["pieces"]).value, 1) or 1; weight = ws.cell(r, ci.get("weight",11)).value
        if kind=="freight" and (skip_freight or norm(ws.cell(r,ci.get("remark",14)).value)==norm("已退件") or pieces != 1):
            stats["跳过运费"]+=1; continue
        if kind == "per_piece_fixed": quote = num(cfg.get("quote", {}).get(key), 0) * pieces
        elif kind == "fixed": quote = num(cfg.get("quote", {}).get(key), 0)
        elif kind == "quote_piece": quote = qb.lookup(channel, name, zone, weight); quote = quote * pieces if quote is not None else None
        elif kind == "freight":
            # SOP 运费1：有“运费更正手续费”时跳过该运费行
            quote = qb.lookup(channel, name, zone, weight)
        else: quote = qb.lookup(channel, name, zone, weight)
        if quote is None:
            stats["未匹配报价"] += 1
            reason="缺少分区" if kind in ("quote","quote_piece","freight") and num(zone) is None else f"渠道/费用未匹配:{channel}/{name}"
            stats["未匹配原因"][reason]=stats["未匹配原因"].get(reason,0)+1
            if reason=="缺少分区": stats["缺少分区"]+=1
        else:
            ws.cell(r, ci["quote"]).value = round(quote, 2); ws.cell(r, ci["diff"]).value = round(quote - amount, 2); stats["已报价"] += 1
            if abs(quote - amount) > cfg.get("run", {}).get("amount_tolerance", .05): stats["差异超容差"] += 1
        if name in ("住宅旺季附加费", "不规则包装") and no:
            stats["强制登记"].append({"no": no, "fee": fee, "amount": amount})
        if mark and no:
            for one_mark in (mark if isinstance(mark,list) else [mark]):
                tasks.append({"row": r, "no": no, "cat": name, "fee": fee, "mark": one_mark, "amount": amount, "note": "按中盟SOP天图核验"})
    # 兼容账单工作簿内嵌的数据列表；单独导出的数据列表由 main_flow 调用同一规则处理。
    for dws in wb.worksheets:
        if "数据列表" not in dws.title: continue
        fee_col=_header_col(dws,("费用名","费用名称"),None)
        channel_col=_header_col(dws,("渠道","渠道名称"),1)
        zone_col=_header_col(dws,("分区","区域"),5)
        weight_col=_header_col(dws,("计费重","单件计费重"),8)
        amount_col=_header_col(dws,("金额","费用金额"),11)
        for r in range(2, dws.max_row + 1):
            if fee_col and norm(dws.cell(r,fee_col).value) != norm("运费"): continue
            channel, zone, weight = dws.cell(r, channel_col).value, dws.cell(r, zone_col).value, dws.cell(r, weight_col).value
            if channel in (None, "") or num(weight) is None: continue
            rate = qb.lookup(channel, "运费", zone, weight)
            if rate is not None:
                dws.cell(r, 21).value = rate
                k = num(dws.cell(r, amount_col).value)
                if k is not None: dws.cell(r, 22).value = round(rate-k, 2)
    stats["天图任务"] = len(tasks)
    if not dry_run: wb.save(path)
    return stats, tasks, bkp

def load_zone_index(path):
    """从中盟“数据列表”导出建立运单号→分区索引（SOP：C列查找、E列分区）。"""
    if not path: return {}
    wb=load_workbook(path,read_only=True,data_only=True); out={}
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2,values_only=True):
            if len(row)>=5 and row[2] not in (None,"") and num(row[4]) is not None:
                out[waybill(row[2])]=int(num(row[4]))
    return out

def write_report(path, stat, tasks, checklist=None):
    """写一份可追溯的运行记录，供人工复核和通知引用。"""
    p=Path(path); p.parent.mkdir(parents=True, exist_ok=True)
    lines=[f"中盟账单数据整理报告 {datetime.now():%Y-%m-%d %H:%M:%S}", f"总行数：{stat['总行数']}", f"已报价：{stat['已报价']}", f"未匹配报价：{stat['未匹配报价']}", f"其中缺少分区：{stat.get('缺少分区',0)}", f"按SOP跳过运费：{stat.get('跳过运费',0)}", f"差异超容差：{stat['差异超容差']}", f"天图任务：{len(tasks)}", f"天图唯一查询组合：{len({(t['no'],t['mark']) for t in tasks})}", f"天图清单：{checklist or '无'}"]
    p.write_text("\n".join(lines)+"\n", encoding="utf-8"); return p

def write_checklist(tasks, out_dir):
    out_dir = Path(out_dir); out_dir.mkdir(parents=True, exist_ok=True); p = out_dir / f"中盟_天图核验清单_{datetime.now():%Y%m%d_%H%M%S}.csv"
    fields = ["Excel行号","客户单号(去后缀)","费用类别","费用名称","期望标记","金额USD","说明"]
    with p.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
        for t in tasks: w.writerow({"Excel行号":t["row"],"客户单号(去后缀)":t["no"],"费用类别":t["cat"],"费用名称":t["fee"],"期望标记":t["mark"],"金额USD":t["amount"],"说明":t["note"]})
    return p
