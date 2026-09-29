# -*- coding: utf-8 -*-
"""加工 FedEx-Details / UPS-Details sheet：直接修改原表。

SOP 步骤（FedEx-Details）：
  1) 在 I 列右侧插入 2 列：J=报价表, K=差异
  2) 在"Original Customer Reference"右侧插入 1 列"费用类型"
     （按 Master 的 C/E 两列做 INDEX/MATCH 查找）——人工模板里该列落在 AY，
     即插在参考号段第一列之后、运单号列之前
  3) 运费：按 (Rated Weight, Zone Code) 查报价表
  4) 百磅：**不设运单数门槛**，明细里所有运单都进百磅透视页

SOP 步骤（UPS-Details）：
  1) 在 N 列右侧插入 4 列：O=费用(中文), P=费用类型, Q=报价表, R=差异
  2) 运费：按 Billed Weight × Zone 查报价
  3) 百磅：只取 O 列翻译为"百磅计费运费"的行进百磅透视页

两个 Details 页**都不再往"报价差异"表写行**（SOP 更新后：百磅差异由
run_skye_check.write_pivot_sheets() 新建的透视页承载，其余费用全在 Master 核对）。
"""

from __future__ import annotations

import logging
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from quote_engine import (
    QuoteBook, _to_float, _to_int, _col_letter_to_index, _index_to_col_letter,
)


@dataclass
class DetailRow:
    row: int
    sheet: str
    tracking: str
    fee_name: str
    amount: float | None
    quote: float | None
    diff: float | None
    note: str = ""


@dataclass
class PivotRow:
    """百磅透视页的一行（一个运单一行）。"""

    tracking: str
    service_code: str
    zone: int | None
    weight: float
    amount: float
    count: int
    rate: float | None = None        # 百磅档位单价，写公式 =rate*总重 用
    min_price: float | None = None   # 保底价，写公式 =MAX(...) 用
    discount: float = 1.0            # 百磅折扣，写公式用
    quote: float | None = None
    diff: float | None = None
    tier: str = ""
    note: str = ""


# ============== Master E→C 索引 ==============

def _resolve_col(header_vals, name: str | None, fallback_letter: str) -> int:
    """优先按表头名定位。

    首次加工会在金额列右侧插列，之后所有列右移，config 里的固定字母随即失效
    （重跑时按字母读会读到空列，导致差异全部归零）。故一律按表头名定位，
    找不到时才退回 config 字母。
    """
    if name:
        for i, h in enumerate(header_vals, 1):
            if h is not None and str(h).strip() == name:
                return i
    return _col_letter_to_index(fallback_letter)


def _hwt_accept(bill: float, info: dict, tolerance: float) -> tuple[float, bool]:
    """百磅有两种计费并存：打折价 与 公布价。任一对上都算无误。

    报价表备注："取消后交单，将按官方账单补收，无法享受折扣，注：会高于报价"、
    "官方复核重量 <200lb 或 >=2000lb，整票会按 UPS Ground 公布价计费"——
    账单本身区分不出是哪种，故两者都接受，只报都不对上的。
    """
    cands = [info["final"]]
    lf = info.get("list_final")
    if lf is not None and abs(lf - info["final"]) > 1e-9:
        cands.append(lf)
    best = min(cands, key=lambda q: abs(q - bill))
    # 按"显示的差异"（2 位小数）判定，否则浮点噪声（-0.20000000000002）
    # 会让差异正好等于容差的行也打上 ⚠，与旁边印出来的数字自相矛盾
    return best, abs(round(best - bill, 2)) <= tolerance


def build_master_index(in_path, sheet_name="Master", key_name="Reference No", val_name="Service Code") -> dict[str, str]:
    """读 Master，按表头名称定位列，建立 Reference No → Service Code 索引。"""
    wb = openpyxl.load_workbook(in_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[sheet_name]
    headers = [str(h).strip() if h else "" for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    try:
        ki = headers.index(key_name)
    except ValueError:
        wb.close()
        return {}
    try:
        vi = headers.index(val_name)
    except ValueError:
        wb.close()
        return {}
    idx: dict[str, str] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or ki >= len(r) or r[ki] is None:
            continue
        k = str(r[ki]).split("-")[0].strip()
        if k and k not in idx and vi < len(r) and r[vi] is not None:
            idx[k] = str(r[vi]).strip()
    wb.close()
    return idx


def _pick_zone(agg: dict, raw_rows: list[list], a_track: int, a_zone: int) -> int | None:
    """取该运单的分区（明细里同一运单分区一致，取第一个非空）。"""
    if agg.get("zone"):
        return agg["zone"]
    for row in raw_rows[1:]:
        if a_track - 1 < len(row) and row[a_track - 1]:
            if str(row[a_track - 1]).split("-")[0].strip() == agg["key"]:
                z = _to_int(row[a_zone - 1] if a_zone - 1 < len(row) else None)
                if z:
                    return z
    return None


# ============== 处理 FedEx-Details ==============

def process_fedex_details(
    in_path, cfg, quote: QuoteBook, log: logging.Logger,
) -> tuple[list[DetailRow], dict]:
    in_path = Path(in_path)
    wb = openpyxl.load_workbook(in_path)
    sheet_name = "FedEx-Details"
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise KeyError(f"账单缺少 {sheet_name} sheet")
    ws = wb[sheet_name]
    dcfg = cfg["details"][sheet_name]
    tolerance = cfg["run"]["amount_tolerance"]

    # 先读原始数据（insert_cols 会导致后续列右移，提前读避免错位）
    raw_rows: list[list] = []
    for row in ws.iter_rows(min_row=1, values_only=False):
        raw_rows.append([cell.value for cell in row])

    # 读取列一律按表头名定位（重跑时列已右移，config 字母会失效）
    hdr0 = raw_rows[0] if raw_rows else []
    a_amount = _resolve_col(hdr0, dcfg.get("amount_header"), dcfg["amount_col"])
    a_weight = _resolve_col(hdr0, dcfg.get("weight_header"), dcfg["weight_col"])
    a_zone   = _resolve_col(hdr0, dcfg.get("zone_header"), dcfg["zone_col"])
    a_track  = _resolve_col(hdr0, dcfg.get("tracking_header"), dcfg["tracking_col"])

    # "费用类型"插在参考号段第一列（Original Customer Reference）右侧，
    # 最终落在 AY 列（与人工模板一致）；取不到该表头时退回运单号列右侧。
    a_feetype = 0
    anchor_name = dcfg.get("feetype_anchor_header")
    if anchor_name:
        for i, h in enumerate(hdr0, 1):
            if h is not None and str(h).strip() == anchor_name:
                a_feetype = i
                break
    if not a_feetype:
        a_feetype = a_track

    has_quote = any(h == "报价表" for h in (raw_rows[0] if raw_rows else []))
    has_feetype = any(h == "费用类型" for h in (raw_rows[0] if raw_rows else []))

    # 按插入点降序处理，避免索引错位
    insertions = []
    if not has_feetype:
        insertions.append((a_feetype + 1, 1, ["费用类型"]))
    if not has_quote:
        insertions.append((a_amount + 1, 2, ["报价表", "差异"]))
    for col_idx, amount, headers in sorted(insertions, key=lambda x: x[0], reverse=True):
        ws.insert_cols(col_idx, amount=amount)
        for i, h in enumerate(headers):
            ws.cell(row=1, column=col_idx + i).value = h

    # 重新用表头名称定位列（不受 insert 影响）
    def _col_by_name(name: str) -> int:
        for cell in ws[1]:
            if cell.value == name:
                return cell.column
        return 1

    q_col = _col_by_name("报价表")
    d_col = _col_by_name("差异")
    y_col = _col_by_name("费用类型")
    i_col = _col_by_name("Transportation Charge Amount")
    t_col = _col_by_name("Rated Weight Amount")
    bk_col = _col_by_name("Zone Code")
    ay_col = _col_by_name("Original Department Reference Description")

    master_idx = build_master_index(in_path)
    log.info("[FedEx] Master 索引 %d 条", len(master_idx))

    problems: list[DetailRow] = []
    agg: dict[str, dict] = {}
    priced: set[int] = set()

    for r in range(2, len(raw_rows) + 1):
        row = raw_rows[r - 1]
        tracking = row[a_track - 1] if a_track - 1 < len(row) else None
        if tracking is None:
            continue
        tracking = str(tracking).split("-")[0].strip()
        amount = _to_float(row[a_amount - 1] if a_amount - 1 < len(row) else None)
        weight = _to_float(row[a_weight - 1] if a_weight - 1 < len(row) else None)
        zone   = _to_int(row[a_zone - 1] if a_zone - 1 < len(row) else None)

        sc = master_idx.get(tracking, "")
        # 费用类型写成公式（SOP：拿 Master 的 C/E 两列当查找区域），与 Python 解析同源
        if sc:
            ws.cell(row=r, column=y_col).value = (
                f'=IFERROR(INDEX(Master!$C:$C,MATCH(LEFT({_index_to_col_letter(ay_col)}{r},'
                f'FIND("-",{_index_to_col_letter(ay_col)}{r}&"-")-1)&"*",Master!$E:$E,0)),"")'
            )
        if not sc:
            continue

        match = quote.service_resolver(sc)

        # FedEx 明细全部参与百磅透视（SOP 更新后不再按"运单数>4"筛）
        g = agg.setdefault(tracking, {
            "key": tracking, "amount": 0.0, "weight": 0.0, "count": 0,
            "zone": None, "sc": sc,
        })
        g["amount"] += amount or 0
        g["weight"] += weight or 0
        g["count"] += 1
        if zone and not g["zone"]:
            g["zone"] = zone

        # 只给 weight_zone 服务（如 FedEx Home Delivery）逐行比 Ground 阶梯价。
        # 百磅服务的行不能逐行套 Ground 价：Multiweigh/MWT 是**多件合并计费**，逐行单价是
        # 折扣后的价（本账单 383 行会全变假差异，人工模板也只给凑巧对上的 8 行写了公式）。
        if not amount or not match or match.kind != "weight_zone":
            continue
        qp = quote.base_price(match.sheet, match.block, weight, zone, round_up=True)
        if qp is None:
            continue
        q_letter = _index_to_col_letter(q_col)
        i_letter = _index_to_col_letter(i_col)
        # 报价表列保留跨工作簿 VLOOKUP（同人工模板 '=VLOOKUP(V2,\'…\]FedEx Home Delivery  \'!$A:$H,8,0)'），
        # 表结构/重量对不上时退回数值，保证这列不会空着。
        formula = quote.vlookup_formula(match.sheet, match.block, r,
                                        _index_to_col_letter(t_col), zone, weight=weight)
        ws.cell(row=r, column=q_col).value = formula or round(qp, 2)
        ws.cell(row=r, column=d_col).value = f"={q_letter}{r}-{i_letter}{r}"
        priced.add(r)
        if abs(round((amount or 0) - qp, 2)) > tolerance:
            problems.append(DetailRow(
                row=r, sheet=sheet_name, tracking=tracking, fee_name="运费",
                amount=amount, quote=qp, diff=round(qp - (amount or 0), 2),
                note=f"运费 R={qp} 账单={amount} 差异={round(qp-(amount or 0),2):+.2f}",
            ))

    # 幂等：这次没给报价的行要擦掉上一轮留下的公式/数值，否则同一张账单重跑结果会不一样
    wiped = 0
    for r in range(2, len(raw_rows) + 1):
        if r in priced:
            continue
        row = raw_rows[r - 1]
        if not (row[a_track - 1] if a_track - 1 < len(row) else None):
            continue
        if ws.cell(row=r, column=q_col).value is not None or \
                ws.cell(row=r, column=d_col).value is not None:
            ws.cell(row=r, column=q_col).value = None
            ws.cell(row=r, column=d_col).value = None
            wiped += 1
    if wiped:
        log.info("[FedEx] 擦掉 %d 行上一轮留下的报价/差异（本轮不报价）", wiped)

    # ---- 百磅透视数据（不写差异表，交给 run_skye_check 新建透视页）----
    pivot_rows: list[PivotRow] = []
    log.info("[FedEx] 百磅透视：%d 单号（全部明细）", len(agg))
    for tracking, g in agg.items():
        sc = g["sc"]
        match = quote.service_resolver(sc)
        zone = _pick_zone(g, raw_rows, a_track, a_zone)
        if not match or match.kind != "hundredweight":
            pivot_rows.append(PivotRow(
                tracking=tracking, service_code=sc, zone=zone,
                weight=g["weight"], amount=round(g["amount"], 2), count=g["count"],
                note=f"非百磅服务（{sc}），运费已在明细列核对，此处不套百磅价",
            ))
            continue
        if zone is None:
            pivot_rows.append(PivotRow(
                tracking=tracking, service_code=sc, zone=None,
                weight=g["weight"], amount=round(g["amount"], 2), count=g["count"],
                note="明细里取不到分区，无法定百磅档位",
            ))
            continue
        info = quote.hwt_price(match.sheet, match.block, g["weight"], zone)
        if not info:
            pivot_rows.append(PivotRow(
                tracking=tracking, service_code=sc, zone=zone,
                weight=g["weight"], amount=round(g["amount"], 2), count=g["count"],
                note=f"报价表 {match.sheet} 匹配不到该重量/分区",
            ))
            continue
        bill = g["amount"]
        final, ok = _hwt_accept(bill, info, tolerance)
        lf = info.get("list_final")
        note = f"{match.sheet} 档={info['tier']}"
        if lf is not None and abs(lf - info["final"]) > 1e-9:
            note += f" 公布价={lf}"
        if not ok:
            note += f" ⚠差异 {round(final-bill,2):+.2f}"
            problems.append(DetailRow(
                row=-1, sheet=sheet_name, tracking=tracking, fee_name="百磅计费",
                amount=round(bill, 2), quote=final, diff=round(final - bill, 2),
                note=(f"百磅 R={final} 账单={bill:.2f} 总重={g['weight']} "
                      f"档={info['tier']} 件数={g['count']} 倍数={bill/final:.4f} "
                      f"差异={round(final-bill,2):+.2f}"),
            ))
        pivot_rows.append(PivotRow(
            tracking=tracking, service_code=sc, zone=zone,
            weight=g["weight"], amount=round(bill, 2), count=g["count"],
            rate=info["rate"], min_price=info["min_price"],
            discount=info.get("discount", 1.0),
            quote=final, diff=round(final - bill, 2),
            tier=info["tier"], note=note,
        ))

    wb.save(in_path)
    wb.close()
    log.info("[FedEx] 已直接写入原表，百磅透视 %d 行", len(pivot_rows))
    return problems, {
        "problems": problems,
        "pivot": {
            "sheet": sheet_name,
            "title": cfg.get("pivot", {}).get(sheet_name, "FedEx-百磅透视"),
            "rows": pivot_rows,
            "meta": {
                "sheet": sheet_name,
                "track_col": _index_to_col_letter(_col_by_name("Original Department Reference Description")),
                "weight_col": _index_to_col_letter(_col_by_name("Rated Weight Amount")),
                "amount_col": _index_to_col_letter(_col_by_name("Transportation Charge Amount")),
                "fee_col": None,
                "fee_value": None,
            },
        },
    }


# ============== 处理 UPS-Details ==============

def process_ups_details(
    in_path, cfg, quote: QuoteBook, translation: dict, log: logging.Logger,
) -> tuple[list[DetailRow], dict]:
    in_path = Path(in_path)
    wb = openpyxl.load_workbook(in_path)
    sheet_name = "UPS-Details"
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise KeyError(f"账单缺少 {sheet_name} sheet")
    ws = wb[sheet_name]
    dcfg = cfg["details"][sheet_name]
    tolerance = cfg["run"]["amount_tolerance"]

    raw_rows: list[list] = []
    for row in ws.iter_rows(min_row=1, values_only=False):
        raw_rows.append([cell.value for cell in row])

    # 读取列按表头名定位（重跑时已插入 O..R 四列，config 字母会失效）
    hdr0 = raw_rows[0] if raw_rows else []
    a_amount = _resolve_col(hdr0, dcfg.get("amount_header"), dcfg["amount_col"])
    a_weight = _resolve_col(hdr0, dcfg.get("weight_header"), dcfg["weight_col"])
    a_zone   = _resolve_col(hdr0, dcfg.get("zone_header"), dcfg["zone_col"])
    a_track  = _resolve_col(hdr0, dcfg.get("tracking_header"), dcfg["tracking_col"])
    a_desc   = _resolve_col(hdr0, dcfg.get("desc_header"), dcfg.get("desc_col", "N"))

    has_fee = any(h == "费用" for h in (raw_rows[0] if raw_rows else []))

    if not has_fee:
        ws.insert_cols(a_desc + 1, amount=4)
        ws.cell(row=1, column=a_desc + 1).value = "费用"
        ws.cell(row=1, column=a_desc + 2).value = "费用类型"
        ws.cell(row=1, column=a_desc + 3).value = "报价表"
        ws.cell(row=1, column=a_desc + 4).value = "差异"

    def _col_by_name(name: str) -> int:
        for cell in ws[1]:
            if cell.value == name:
                return cell.column
        return 1

    o_col = _col_by_name("费用")
    p_col = _col_by_name("费用类型")
    q_col = _col_by_name("报价表")
    r_col = _col_by_name("差异")
    s_col = _col_by_name("Net Amount")
    i_col = _col_by_name("Billed Weight")
    l_col = _col_by_name("Zone")
    g_col = _col_by_name("Shipment Reference Number 1")
    n_col = _col_by_name("Charge Description")

    master_idx = build_master_index(in_path)
    log.info("[UPS] Master 索引 %d 条", len(master_idx))

    problems: list[DetailRow] = []
    agg: dict[str, dict] = {}
    priced: set[int] = set()

    for r in range(2, len(raw_rows) + 1):
        row = raw_rows[r - 1]
        desc = row[a_desc - 1] if a_desc - 1 < len(row) else None
        tracking = row[a_track - 1] if a_track - 1 < len(row) else None
        if tracking is None:
            continue
        tracking = str(tracking).split("-")[0].strip()
        amount = _to_float(row[a_amount - 1] if a_amount - 1 < len(row) else None)
        weight = _to_float(row[a_weight - 1] if a_weight - 1 < len(row) else None)
        zone   = _to_int(row[a_zone - 1] if a_zone - 1 < len(row) else None)

        desc_str = str(desc).strip() if desc else ""
        ws.cell(row=r, column=o_col).value = translation.get(desc_str, desc_str)
        sc = master_idx.get(tracking, "")
        g_letter = _index_to_col_letter(g_col)
        if sc:
            ws.cell(row=r, column=p_col).value = (
                f'=IFERROR(INDEX(Master!$C:$C,MATCH(LEFT({g_letter}{r},'
                f'FIND("-",{g_letter}{r}&"-")-1)&"*",Master!$E:$E,0)),"")'
            )
        if not sc or not amount:
            continue

        fee_zh = translation.get(desc_str, desc_str)
        if fee_zh not in ("地面商业-运费", "地面住宅-运费", "百磅计费运费"):
            continue

        match = quote.service_resolver(sc)
        if not match:
            continue
        if fee_zh == "百磅计费运费":
            # SOP 054：按 O 列中文翻译筛出"百磅计费运费"，不做运单数门槛，全部处理
            g = agg.setdefault(tracking, {
                "key": tracking, "amount": 0.0, "weight": 0.0, "count": 0,
                "zone": None, "sc": sc,
            })
            g["amount"] += amount
            g["weight"] += weight or 0
            g["count"] += 1
            if zone and not g["zone"]:
                g["zone"] = zone
        elif match.kind == "weight_zone":
            qp = quote.base_price(match.sheet, match.block, weight, zone, round_up=True)
            if qp is not None:
                # 报价表列保留跨工作簿 VLOOKUP（同人工模板 '=VLOOKUP(I4,\'…\]UPS-Ground商业\'!$B:$I,8,0)'）
                formula = quote.vlookup_formula(match.sheet, match.block, r,
                                                _index_to_col_letter(i_col), zone,
                                                weight=weight)
                ws.cell(row=r, column=q_col).value = formula or round(qp, 2)
                ws.cell(row=r, column=r_col).value = f"={_index_to_col_letter(q_col)}{r}-{_index_to_col_letter(s_col)}{r}"
                priced.add(r)
                if abs(round(amount - qp, 2)) > tolerance:
                    problems.append(DetailRow(
                        row=r, sheet=sheet_name, tracking=tracking, fee_name=fee_zh,
                        amount=amount, quote=qp, diff=round(qp - amount, 2),
                        note=f"运费 R={qp} 账单={amount} 差异={round(qp-amount,2):+.2f}",
                    ))

    # 幂等：这次没给报价的行要擦掉上一轮留下的公式/数值
    wiped = 0
    for r in range(2, len(raw_rows) + 1):
        if r in priced:
            continue
        row = raw_rows[r - 1]
        if not (row[a_track - 1] if a_track - 1 < len(row) else None):
            continue
        if ws.cell(row=r, column=q_col).value is not None or \
                ws.cell(row=r, column=r_col).value is not None:
            ws.cell(row=r, column=q_col).value = None
            ws.cell(row=r, column=r_col).value = None
            wiped += 1
    if wiped:
        log.info("[UPS] 擦掉 %d 行上一轮留下的报价/差异（本轮不报价）", wiped)

    pivot_rows: list[PivotRow] = []
    log.info("[UPS] 百磅透视 %d 单号", len(agg))
    for tracking, g in agg.items():
        sc = g["sc"]
        match = quote.service_resolver(sc)
        zone = _pick_zone(g, raw_rows, a_track, a_zone)
        if not match or match.kind != "hundredweight":
            pivot_rows.append(PivotRow(
                tracking=tracking, service_code=sc, zone=zone,
                weight=g["weight"], amount=round(g["amount"], 2), count=g["count"],
                note=f"非百磅服务（{sc}）",
            ))
            continue
        if zone is None:
            pivot_rows.append(PivotRow(
                tracking=tracking, service_code=sc, zone=None,
                weight=g["weight"], amount=round(g["amount"], 2), count=g["count"],
                note="明细里取不到分区，无法定百磅档位",
            ))
            continue
        info = quote.hwt_price(match.sheet, match.block, g["weight"], zone)
        if not info:
            pivot_rows.append(PivotRow(
                tracking=tracking, service_code=sc, zone=zone,
                weight=g["weight"], amount=round(g["amount"], 2), count=g["count"],
                note=f"报价表 {match.sheet} 匹配不到该重量/分区",
            ))
            continue
        bill = g["amount"]
        final, ok = _hwt_accept(bill, info, tolerance)
        lf = info.get("list_final")
        note = f"{match.sheet} 档={info['tier']}"
        if lf is not None and abs(lf - info["final"]) > 1e-9:
            note += f" 公布价={lf}"
        if not ok:
            note += f" ⚠差异 {round(final-bill,2):+.2f}"
            problems.append(DetailRow(
                row=-1, sheet=sheet_name, tracking=tracking, fee_name="百磅计费",
                amount=round(bill, 2), quote=final, diff=round(final - bill, 2),
                note=(f"百磅 R={final} 账单={bill:.2f} 总重={g['weight']} "
                      f"档={info['tier']} 行数={g['count']} 倍数={bill/final:.4f} "
                      f"差异={round(final-bill,2):+.2f}"),
            ))
        pivot_rows.append(PivotRow(
            tracking=tracking, service_code=sc, zone=zone,
            weight=g["weight"], amount=round(bill, 2), count=g["count"],
            rate=info["rate"], min_price=info["min_price"],
            discount=info.get("discount", 1.0),
            quote=final, diff=round(final - bill, 2),
            tier=info["tier"], note=note,
        ))

    wb.save(in_path)
    wb.close()
    log.info("[UPS] 已直接写入原表，百磅透视 %d 行", len(pivot_rows))
    return problems, {
        "problems": problems,
        "pivot": {
            "sheet": sheet_name,
            "title": cfg.get("pivot", {}).get(sheet_name, "UPS-百磅透视"),
            "rows": pivot_rows,
            "meta": {
                "sheet": sheet_name,
                "track_col": _index_to_col_letter(g_col),
                "weight_col": _index_to_col_letter(i_col),
                "amount_col": _index_to_col_letter(s_col),
                "fee_col": _index_to_col_letter(o_col),
                "fee_value": "百磅计费运费",
            },
        },
    }


# ============== 顶层入口 ==============

def process_details(
    in_path, cfg, quote: QuoteBook, translation: dict, log: logging.Logger,
) -> dict:
    _, f = process_fedex_details(in_path, cfg, quote, log)
    _, u = process_ups_details(in_path, cfg, quote, translation, log)
    return {"fedex": f, "ups": u}
