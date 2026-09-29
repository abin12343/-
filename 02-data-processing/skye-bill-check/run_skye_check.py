# -*- coding: utf-8 -*-
"""SKYE 账单核价引擎（CLI）。

直接修改原表，输出报告 + 天图核验清单到 output_dir。
  python run_skye_check.py --input <账单.xlsx> [--dry-run] [--skip-master|--skip-details]

可作为子进程被 main_flow.py 拉起，也可单独跑。
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from quote_engine import QuoteBook
from process_master import process_master, load_translation, PriceRow
from process_details import process_details


if getattr(sys, "frozen", False):
    HERE = Path(sys.executable).resolve().parent.parent
    # 冻结后 PYTHONUTF8 不生效，管道默认按 cp936 写中文，父流程却按 utf-8 解码。
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
else:
    HERE = Path(__file__).resolve().parent


def _setup_log(name: str = "skye-check") -> logging.Logger:
    log = logging.getLogger(name)
    if not log.handlers:
        log.setLevel(logging.INFO)
        fmt = logging.Formatter("[%(asctime)s] %(message)s", datefmt="%H:%M:%S")
        ch = logging.StreamHandler(sys.stdout)
        ch.setFormatter(fmt)
        log.addHandler(ch)
    return log


def load_config() -> dict:
    cfg_file = HERE / "config.json"
    if not cfg_file.is_file():
        raise SystemExit(f"找不到 config.json：{cfg_file}")
    return json.loads(cfg_file.read_text(encoding="utf-8"))


def _pv(p, *names, default=None):
    """按候选属性名取值（PriceRow 与 DetailRow 字段名不同）。"""
    for n in names:
        v = getattr(p, n, None)
        if v not in (None, ""):
            return v
    return default


def write_price_diff_sheet(wb, problems: list):
    """在原表里追加/覆盖"报价差异"sheet。

    只收 Master 的差异（SOP 更新后：FedEx-Details / UPS-Details 的百磅差异
    改由 write_pivot_sheets() 新建的透视页承载，不再写这张表）。
    """
    if "报价差异" in wb.sheetnames:
        del wb["报价差异"]
    ws = wb.create_sheet("报价差异")
    # SOP 034-036：要能看出 运单号 / 费用名称 / 金额 / 分区
    headers = ["Excel行号", "Sheet", "客户单号", "服务代码", "费用名称",
               "账单金额", "报价", "差异", "分区", "备注"]
    ws.append(headers)
    for p in problems:
        ws.append([
            p.row, p.sheet,
            _pv(p, "reference_no", "tracking", default=""),
            _pv(p, "service_code", default=""),
            p.fee_name,
            _pv(p, "amount"), _pv(p, "quote"), _pv(p, "diff"),
            _pv(p, "zone"),
            p.note,
        ])
    return ws


# SOP 018-035：两张百磅透视页的列（严格按 SOP 口径，多余列一律不要）
#   FedEx-Details：I=金额 V=重量 AX=运单号(引用) AY=费用类型 BN=分区
#                  → 透视金额/重量/运单号，再引用运单号 匹配**费用类型和分区**
#   UPS-Details  ：G=运单号 I=重量 L=分区 S=金额
#                  → 透视重量/金额/运单号，再引用运单号 匹配**分区**
PIVOT_COLS = {
    "fedex": ["运单号", "费用类型", "分区", "总重量", "账单金额", "报价", "差异"],
    "ups":   ["运单号", "分区", "总重量", "账单金额", "报价", "差异"],
}


def _sum_formula(meta: dict, key_cell: str, val_col: str) -> str:
    """SUMIF/SUMIFS：按运单号聚合明细列。

    运单号用 `key&"*"` 通配，明细里带 `-2` 后缀的行会一起并进来
    （和 Python 侧 split('-')[0] 的口径一致）。
    """
    sh = meta["sheet"]
    trk = meta["track_col"]
    if meta.get("fee_col"):
        # UPS：只聚合"百磅计费运费"那些行
        return (f"=SUMIFS('{sh}'!${val_col}:${val_col},"
                f"'{sh}'!${trk}:${trk},{key_cell}&\"*\","
                f"'{sh}'!${meta['fee_col']}:${meta['fee_col']},\"{meta['fee_value']}\")")
    return (f"=SUMIF('{sh}'!${trk}:${trk},{key_cell}&\"*\","
            f"'{sh}'!${val_col}:${val_col})")


def _round(x: float) -> str:
    """写进公式的数字：去掉多余小数位，避免 0.38020000000000004 这种。"""
    return f"{round(float(x), 6):g}"


def write_pivot_sheets(wb, res: dict, log: logging.Logger):
    """新建两张百磅透视页（严格按 SOP 018-035 搭）。

    整页公式：总重量/账单金额用 SUMIF(S) 回明细页取（等价于 SOP 的"粘贴 I/V/AX、G/I/L/S
    再透视"，openpyxl 建不了真正的 Excel 数据透视表，用 SUMIF 得到同样的数），
    报价 = 该百磅档位单价 × 总重量，差异 = 报价 − 账单（SOP 说的"加减法"）。
    改重量或改单价，结果自动跟着变。

    :return: (建了几张透视页, 无法定价的行说明)。说明只进报告 txt——
        透视页本身严格只留 SOP 那几列，不塞备注。
    """
    made = 0
    notes: list[str] = []
    for key in ("fedex", "ups"):
        node = (res or {}).get(key) or {}
        pv = node.get("pivot") or {}
        rows = pv.get("rows") or []
        title = pv.get("title")
        meta = pv.get("meta") or {}
        if not title or not meta.get("track_col"):
            continue
        headers = PIVOT_COLS[key]
        col = {h: get_column_letter(i) for i, h in enumerate(headers, start=1)}
        if title in wb.sheetnames:
            del wb[title]
        ws = wb.create_sheet(title)
        ws.append(headers)
        w = col["总重量"]
        b = col["账单金额"]
        q = col["报价"]
        for i, p in enumerate(rows, start=2):
            quote_f = diff_f = None
            if p.rate is not None:
                rate = _round(p.rate)
                # 百磅报价严格按 SOP：报价表单价 × 该运单汇总总重量。
                # 不再叠加折扣或最低收费，避免报价金额被人为放大/缩小。
                quote_f = f"=ROUND({rate}*{w}{i},2)"
                # 差异统一使用最直观的口径：报价 - 账单金额。
                # 报价列已经给出本行采用的报价，差异列只引用这两列，便于使用者理解。
                diff_f = f"={q}{i}-{b}{i}"
            elif p.note:
                notes.append(f"{title} 行{i} 运单{p.tracking}：{p.note}")
            vals = {
                "运单号": p.tracking,
                "费用类型": p.service_code,
                "分区": p.zone,
                "总重量": _sum_formula(meta, f"A{i}", meta["weight_col"]),
                "账单金额": _sum_formula(meta, f"A{i}", meta["amount_col"]),
                "报价": quote_f,
                "差异": diff_f,
            }
            ws.append([vals[h] for h in headers])
        made += 1
        log.info("[透视] %s %d 行（列：%s）", title, len(rows), "/".join(headers))
    if notes:
        log.info("[透视] 无法定价 %d 行（说明写在报告里，不占透视页的列）", len(notes))
    return made, notes


def _backup_original(in_path: Path, cfg: dict, log: logging.Logger) -> Path | None:
    """加工前把原始账单另存为 <名>_原始备份.xlsx。

    已存在时**不覆盖**：否则反复运行时会把"已加工"的表当成原始表存下来，
    备份就失去意义（这也正是"后续处理默认之前处理已存在"的来源）。
    """
    if not cfg.get("run", {}).get("backup_original", True):
        return None
    bak = in_path.with_name(f"{in_path.stem}_原始备份{in_path.suffix}")
    if bak.exists():
        log.info("[备份] 已存在，保持不变：%s", bak.name)
        return bak
    shutil.copy2(in_path, bak)
    log.info("[备份] 原始账单 -> %s", bak.name)
    return bak


def write_tiantu_checklist(out_path, problems: list[PriceRow], cfg, log):
    """生成天图核验清单 CSV（兼容 永达 check_tiantu.py 的字段）。

    :param problems: **所有被收费的**附加费行（不只核价有差异的）——
        SOP 要求费用行都要去天图验标记。按费用名归到费用类别，
        再按 (类别, 单号) 去重（同一票同一个费只需验一次标记）。
    """
    checks = cfg.get("tiantu_checks", {})
    # 直接按费用名归类，不再借助列字母（列标会随加工插列而变，极易错位）
    fee_to_cat = {
        "超重附加费":     "overweight",
        "超尺寸附加费":   "oversize",
        "超大尺寸费用":   "oversize",
        "地址更正":       "address_correction",
        "商业偏远附加费": "das_commercial",
        "住宅偏远附加费": "das_residential",
        "商业超偏远附加费": "das_extended",
        "私人超偏远":     "das_extended",
        "私人住宅附加费": "residential",
    }
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    rows: list[dict] = []
    seen = set()
    for p in problems:
        if not p.reference_no:
            continue
        cat = fee_to_cat.get(p.fee_name)
        if not cat or cat not in checks:
            continue
        no = p.reference_no.split("-")[0]
        key = (cat, no)
        if key in seen:
            continue
        seen.add(key)
        rows.append({
            "row": p.row,
            "no": no,
            "cat": cat,
            "fee": p.fee_name,
            "mark": checks[cat]["mark"],
            "amount": p.amount,
            "note": p.note,
        })
    with out.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Excel行号", "客户单号(去后缀)", "费用类别", "费用名称", "期望标记", "金额USD", "说明"])
        for r in rows:
            w.writerow([r["row"], r["no"], r["cat"], r["fee"], r["mark"], r["amount"], r["note"]])
    log.info("[checklist] 天图核验任务 %d 条（去重前收费行 %d 行）-> %s",
             len(rows), len(problems or []), out)
    return out, len(rows)


def write_report(out_path, bill_name, freight_probs, surcharge_probs, fedex_probs, ups_probs, stamp,
                 pivot_notes=None):
    lines = [
        f"SKYE 账单核价报告 - {bill_name}",
        f"运行时间：{datetime.now():%Y-%m-%d %H:%M:%S}",
        f"任务编号：{stamp}",
        "",
        f"【Master 运费差异】 {len(freight_probs)} 行",
    ]
    for p in freight_probs[:30]:
        lines.append(f"  行{p.row} {p.service_code} 运单{p.reference_no}  {p.note}")
    if len(freight_probs) > 30:
        lines.append(f"  ... 其余 {len(freight_probs)-30} 行省略")
    lines.append("")
    lines.append(f"【Master 附加费差异/无报价】 {len(surcharge_probs)} 行")
    for p in surcharge_probs[:30]:
        lines.append(f"  行{p.row} {p.service_code} 运单{p.reference_no}  {p.note}")
    if len(surcharge_probs) > 30:
        lines.append(f"  ... 其余 {len(surcharge_probs)-30} 行省略")
    lines.append("")
    lines.append(f"【FedEx-Details 运费差异（不入报价差异表）】 {len(fedex_probs)} 行")
    for p in fedex_probs[:20]:
        lines.append(f"  行{p.row} 运单{p.tracking}  {p.note}")
    if len(fedex_probs) > 20:
        lines.append(f"  ... 其余 {len(fedex_probs)-20} 行省略")
    lines.append("")
    lines.append(f"【UPS-Details 运费差异（不入报价差异表）】 {len(ups_probs)} 行")
    for p in ups_probs[:20]:
        lines.append(f"  行{p.row} 运单{p.tracking}  {p.note}")
    if len(ups_probs) > 20:
        lines.append(f"  ... 其余 {len(ups_probs)-20} 行省略")
    notes = list(pivot_notes or [])
    lines.append("")
    lines.append(f"【百磅透视页无法定价的行（透视页里报价/差异留空）】 {len(notes)} 行")
    for n in notes[:30]:
        lines.append(f"  {n}")
    if len(notes) > 30:
        lines.append(f"  ... 其余 {len(notes)-30} 行省略")
    Path(out_path).write_text("\n".join(lines), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description="SKYE 账单核价引擎")
    ap.add_argument("--input", required=True, help="账单 xlsx 路径")
    ap.add_argument("--quote", default=None,
                    help="报价表 xlsx 路径（报价表时常更换；不传则用 config.paths.quote_file）")
    ap.add_argument("--skip-master", action="store_true")
    ap.add_argument("--skip-details", action="store_true")
    ap.add_argument("--dry-run", action="store_true", help="只跑不写盘")
    ap.add_argument("--no-write-dialog", action="store_true", help="跳过弹窗（用于 main_flow 串联）")
    args = ap.parse_args()

    cfg = load_config()
    log = _setup_log()
    in_path = Path(args.input)
    if not in_path.is_file():
        raise SystemExit(f"账单文件不存在：{in_path}")

    quote_path = Path(args.quote) if args.quote else Path(cfg["paths"]["quote_file"])
    if not quote_path.is_file():
        raise SystemExit(f"报价表不存在：{quote_path}")
    quote = QuoteBook(
        str(quote_path), cfg["quote"]["service_map"],
        hwt_discount=cfg["quote"].get("hundredweight_discount", 1.0),
        discount_min_price=cfg["quote"].get("discount_min_price", False),
        book_dir=in_path.parent,   # 公式里写相对路径，换台电脑也能匹配
    )
    trans = load_translation(cfg["paths"]["translation_file"])
    log.info("翻译表 %d 条，报价表 %s", len(trans), quote_path)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = Path(cfg["paths"]["output_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)

    freight_probs: list[PriceRow] = []
    surcharge_probs: list[PriceRow] = []
    tiantu_rows: list[PriceRow] = []
    cl_count = 0
    fedex_probs = []
    ups_probs = []

    # 加工前备份原始账单（已存在则不覆盖）
    if not args.dry_run:
        _backup_original(in_path, cfg, log)

    # Master
    if not args.skip_master:
        if args.dry_run:
            log.info("[dry-run] 跳过 Master 加工")
        else:
            freight_probs, surcharge_probs, tiantu_rows = process_master(in_path, cfg, quote, trans, log)

    # Details
    details_res: dict = {}
    if not args.skip_details:
        if args.dry_run:
            log.info("[dry-run] 跳过 Details 加工")
        else:
            details_res = process_details(in_path, cfg, quote, trans, log)
            fedex_probs = details_res["fedex"]["problems"]
            ups_probs = details_res["ups"]["problems"]

    # 报价差异 sheet（仅 Master）+ 百磅透视页（FedEx/UPS 各一张）
    pivot_notes: list[str] = []
    if not args.dry_run:
        try:
            wb = openpyxl.load_workbook(in_path)
            write_price_diff_sheet(wb, list(freight_probs) + list(surcharge_probs))
            if details_res:
                _, pivot_notes = write_pivot_sheets(wb, details_res, log)
            wb.save(in_path)
            wb.close()
        except Exception as exc:
            log.warning("追加差异/透视 sheet 失败：%s", exc)

    # 报告
    report_path = out_dir / f"{in_path.stem}_报告_{stamp}.txt"
    if not args.dry_run:
        write_report(report_path, in_path.name, freight_probs, surcharge_probs,
                     fedex_probs, ups_probs, stamp, pivot_notes)
        log.info("报告：%s", report_path)

    # 天图核验清单
    if not args.dry_run:
        cl = out_dir / f"{in_path.stem}_天图核验清单_{stamp}.csv"
        _, cl_count = write_tiantu_checklist(cl, tiantu_rows, cfg, log)

    # 控制台汇总
    print()
    print("=" * 60)
    print(f"账单：{in_path.name}")
    print(f"  Master 运费差异:        {len(freight_probs)} 行")
    print(f"  Master 附加费差异/无报价: {len(surcharge_probs)} 行")
    print(f"  （以上写入「报价差异」表）")
    print(f"  FedEx-Details 运费差异:  {len(fedex_probs)} 行")
    print(f"  UPS-Details 运费差异:    {len(ups_probs)} 行")
    print(f"  （以上只在报告里；百磅差异见「{cfg['pivot']['FedEx-Details']}」「{cfg['pivot']['UPS-Details']}」页）")
    if not args.dry_run:
        print(f"  天图核验清单（全部收费行）: {cl_count} 票/费用")
        if pivot_notes:
            print(f"  百磅透视页无法定价:      {len(pivot_notes)} 行（说明见报告）")
        print(f"  报告/清单：{out_dir}")
    print("=" * 60)


if __name__ == "__main__":
    main()
