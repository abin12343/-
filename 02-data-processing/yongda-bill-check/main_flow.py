# -*- coding: utf-8 -*-
"""
主流程入口：一条命令跑完全程
  1) 弹窗选择要处理的账单 xlsx（不再自动下载压缩包）
  2) run_yongda_check.py（核价 + 四列加工副本 + 生成天图核验清单）
  3) check_tiantu.py（天图逐类核验）
  4) write_problems_to_stats.py（合并核价问题 + 天图缺标记问题，弹窗写入统计表）
结束后自动写金山文档运行记录，并向企业微信群推送结果（推送失败不影响主流程）。

运行：python main_flow.py
常用参数：
  --input <账单xlsx/目录>    跳过弹窗，直接用指定账单
  --skip-engine / --skip-tiantu / --skip-write
  --tiantu-limit N          天图只测前 N 个任务（调试用）
  --dry-run                 最后一步只打印待写内容，不弹窗
"""

from __future__ import annotations

import argparse
import http.client
import json
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path


if getattr(sys, "frozen", False):
    HERE = Path(sys.executable).resolve().parent.parent
else:
    HERE = Path(__file__).resolve().parent
PY = sys.executable

# ========== 配置 ==========

KDOCS_CONFIG = {
    "FILE_ID": "",
    "TOKEN": "",
    "SCRIPT_ID": "",
    "SHEET_NAME": "数据表",
    "FIELD_MAPPING": {
        "app_name": "应用名称",
        "duration": "运行时长",
        "run_time": "运行时间",
        "account": "账号标识",
        "status": "运行状态",
    },
}
RETRY_TIMES = 3
RETRY_DELAY = 2
TIMEOUT = 10
WECOM_WEBHOOK_URL = ""
APP_NAME = "天图财务-永达账单数据整理"
ACCOUNT_ID = "脚本"


def _load_notify_local():
    """从本机 notify.local.json（不入库）读取通知配置；没有该文件时保持空值。"""
    global WECOM_WEBHOOK_URL, APP_NAME, ACCOUNT_ID
    try:
        data = json.loads((HERE / "notify.local.json").read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return
    if data.get("file_id"):
        KDOCS_CONFIG["FILE_ID"] = str(data["file_id"])
    if data.get("token"):
        KDOCS_CONFIG["TOKEN"] = str(data["token"])
    if data.get("script_id"):
        KDOCS_CONFIG["SCRIPT_ID"] = str(data["script_id"])
    if data.get("sheet_name"):
        KDOCS_CONFIG["SHEET_NAME"] = str(data["sheet_name"])
    if data.get("wecom_key"):
        WECOM_WEBHOOK_URL = (
            "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key="
            + str(data["wecom_key"])
        )
    if data.get("app_name"):
        APP_NAME = str(data["app_name"])
    if data.get("account_id"):
        ACCOUNT_ID = str(data["account_id"])


_load_notify_local()


def post_https_json(host, path, payload, headers, timeout=TIMEOUT):
    """标准库 HTTPS POST，返回 (http状态码, 响应文本)。"""
    conn = http.client.HTTPSConnection(host, timeout=timeout)
    try:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        conn.request("POST", path, body=body, headers=headers)
        resp = conn.getresponse()
        text = resp.read().decode("utf-8", "replace")
        return resp.status, text
    finally:
        conn.close()


def write_run_record(app_name, start_ts, account_id, status, config=None):
    """按影刀 SOP 提供的金山文档接口写入本次运行记录；失败只返回消息，不抛异常。"""
    cfg = config or KDOCS_CONFIG
    try:
        try:
            start_ts = float(start_ts)
        except (TypeError, ValueError):
            start_ts = time.time()
        duration = round(time.time() - start_ts, 2)
        run_time_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        field_map = cfg["FIELD_MAPPING"]
        fields = {
            field_map["app_name"]: app_name,
            field_map["duration"]: duration,
            field_map["run_time"]: run_time_str,
            field_map["account"]: account_id,
            field_map["status"]: status,
        }
        record_list = [{"fields": fields}]

        # ====== 关键点：外层 "Context" ======
        payload = {
            "Context": {
                "argv": {
                    "sheet_name": cfg["SHEET_NAME"],
                    "sheet_data": record_list,
                }
            }
        }
        headers = {
            "Content-Type": "application/json",
            "AirScript-Token": cfg["TOKEN"],
        }
        host = "www.kdocs.cn"
        api_path = (
            f"/api/v3/ide/file/{cfg['FILE_ID']}/script/{cfg['SCRIPT_ID']}/sync_task"
        )
        last = ""
        for attempt in range(1, RETRY_TIMES + 1):
            try:
                status_code, text = post_https_json(host, api_path, payload, headers)
                if status_code == 200:
                    return True, f"[OK] 写入成功，应用: {app_name}，耗时: {duration}秒", text
                if 400 <= status_code < 500:
                    return False, f"[失败] 客户端错误 ({status_code}): {text}", text
                last = f"服务端错误 HTTP {status_code}: {text[:200]}"
            except Exception as exc:  # noqa: BLE001
                last = f"{exc}"
            if attempt < RETRY_TIMES:
                time.sleep(RETRY_DELAY)
        return False, f"[失败] 所有重试均失败: {last}", ""
    except Exception as exc:  # noqa: BLE001
        return False, f"[失败] 执行异常: {exc}", ""


def send_wecom_push(content):
    """企业微信群机器人 text 消息推送；失败只返回消息，不抛异常。"""
    payload = {"msgtype": "text", "text": {"content": content}}
    headers = {"Content-Type": "application/json"}
    host = "qyapi.weixin.qq.com"
    path = "/cgi-bin/webhook/send?" + WECOM_WEBHOOK_URL.split("?", 1)[1]
    last = ""
    for attempt in range(1, RETRY_TIMES + 1):
        try:
            status_code, text = post_https_json(host, path, payload, headers)
            if status_code == 200:
                try:
                    data = json.loads(text)
                except Exception:  # noqa: BLE001
                    data = {}
                if data.get("errcode") in (0, None):
                    return True, "[OK] 企业微信推送成功"
                last = f"企业微信返回: {text[:200]}"
            else:
                last = f"HTTP {status_code}: {text[:200]}"
        except Exception as exc:  # noqa: BLE001
            last = f"{exc}"
        if attempt < RETRY_TIMES:
            time.sleep(RETRY_DELAY)
    return False, f"[失败] 企业微信推送失败: {last}"


def build_summary(status, start_ts, stage_reports, note=""):
    lines = [
        f"应用：{APP_NAME}",
        f"状态：{status}",
        f"运行时间：{datetime.now():%Y-%m-%d %H:%M:%S}",
        f"运行时长：{round(time.time() - start_ts, 1)} 秒",
    ]
    if stage_reports:
        parts = []
        for s in stage_reports:
            mark = "完成" if s["ok"] else "失败"
            parts.append(f"{s['name']} {s['secs']}s({mark})")
        lines.append("阶段：" + " | ".join(parts))
    if note:
        lines.append(f"备注：{note}")
    return "\n".join(lines)


def notify_result(status, start_ts, stage_reports, note=""):
    ok1, msg1, _ = write_run_record(APP_NAME, start_ts, ACCOUNT_ID, status)
    ok2, msg2 = send_wecom_push(build_summary(status, start_ts, stage_reports, note))
    _safe_print(f"[通知] 金山文档记录: {msg1}")
    _safe_print(f"[通知] {msg2}")


def _safe_print(msg):
    """控制台编码兼容输出（某些 Windows 控制台不支持 emoji 等字符，避免中断流程）。"""
    try:
        print(msg)
    except UnicodeEncodeError:
        enc = getattr(sys.stdout, "encoding", None) or "utf-8"
        sys.stdout.buffer.write((msg + "\n").encode(enc, "replace"))
        sys.stdout.flush()


def run_stage(name, script, extra=None):
    print("\n" + "=" * 60)
    print(f"[主流程] {name}: {script.name}")
    print("=" * 60)
    t0 = time.time()
    if getattr(sys, "frozen", False):
        exe = HERE / script.stem / f"{script.stem}.exe"
        cmd = [str(exe)] + (extra or [])
    else:
        cmd = [PY, str(script)] + (extra or [])
    code = subprocess.run(cmd).returncode
    secs = round(time.time() - t0, 2)
    print(f"[主流程] {name} 用时 {secs}s（退出码 {code}）")
    return code, secs


def choose_bill_via_dialog():
    """弹窗选择要处理的账单 xlsx；校验不通过会提示并重新弹出，取消返回 None。"""
    from tkinter import messagebox

    initial = ""
    try:
        cfg = json.loads((HERE / "config.json").read_text(encoding="utf-8"))
        p = Path(cfg["paths"]["input_dir"])
        initial = str(p) if p.exists() else ""
    except Exception:  # noqa: BLE001
        pass
    for _ in range(10):
        path = _ask_open_file(
            title="请选择要处理的账单 xlsx（如 TTTX_UPS*.xlsx；同批多份请选第一份）",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")],
            initialdir=initial or str(Path.home()),
        )
        if not path:
            return None
        err = _validate_bill_file(path)
        if err is None:
            return path
        messagebox.showerror("文件不符合要求", err)
    print("已连续多次选择不符合要求的文件，流程取消。")
    return None


def _ask_open_file(title, filetypes, initialdir=None):
    """独立临时 Tk 窗口弹出文件选择；关闭后窗口即销毁，避免隐藏窗口导致后续弹窗卡住。"""
    import tkinter as tk
    from tkinter import filedialog

    root = tk.Tk()
    root.withdraw()
    try:
        return filedialog.askopenfilename(
            title=title,
            filetypes=filetypes,
            initialdir=initialdir or "",
        )
    finally:
        root.destroy()


def _validate_bill_file(path):
    """校验待处理账单 xlsx；通过返回 None，否则返回可展示的错误说明。"""
    p = Path(path or "")
    if not p.exists():
        return "文件不存在，请重新选择。"
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        return f"不是 Excel 工作簿：{p.name}，请选择 .xlsx 文件。"
    try:
        import openpyxl
    except Exception:  # noqa: BLE001
        return None
    wb = None
    try:
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        header = "".join("" if c is None else str(c).strip() for c in row)
    except Exception as exc:  # noqa: BLE001
        return f"无法打开该 Excel 文件：{exc}"
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:  # noqa: BLE001
                pass
    if "费用金额" not in header or not any(k in header for k in ("系统单号", "客户单号")):
        return "选择的文件不是永达账单（缺少 费用金额/系统单号/客户单号 表头），请重新选择。"
    return None


def latest(out_dir: Path, pattern: str):
    cands = sorted(
        out_dir.glob(pattern),
        key=lambda f: f.stat().st_mtime,
        reverse=True,
    )
    return cands[0] if cands else None


def default_output_dir():
    """统一读取 config.json 的输出目录，保证与各子脚本同一处。"""
    try:
        cfg_file = HERE / "config.json"
        cfg = json.loads(cfg_file.read_text(encoding="utf-8"))
        p = Path(cfg["paths"]["output_dir"])
        return p if p.is_absolute() else cfg_file.resolve().parent / p
    except Exception:  # noqa: BLE001
        return HERE / "outputs"


def main():
    ap = argparse.ArgumentParser(description="永达账单对账全流程（弹窗选表）")
    ap.add_argument("--input", default=None, help="账单 xlsx 或目录（提供则跳过选表弹窗）")
    ap.add_argument("--skip-fetch", action="store_true",
                    help="已废弃：不再自动取账单，保留该参数仅为兼容")
    ap.add_argument("--skip-engine", action="store_true")
    ap.add_argument("--skip-tiantu", action="store_true")
    ap.add_argument("--skip-write", action="store_true")
    ap.add_argument("--tiantu-limit", type=int, default=0, help="天图只测前 N 个任务")
    ap.add_argument("--dry-run", action="store_true", help="最后一步只打印，不弹窗不写入")
    args = ap.parse_args()

    out_dir = default_output_dir()

    # 0) 选择要处理的账单（不再自动下载）
    bill_arg = args.input
    if not bill_arg:
        print("\n" + "=" * 60)
        print("[主流程] 请在弹出的窗口中选择要处理的账单 xlsx（不再自动下载压缩包）")
        print("=" * 60)
        bill_arg = choose_bill_via_dialog()
        if not bill_arg:
            print("未选择账单文件，流程取消。")
            return

    start_ts = time.time()
    stage_reports = []
    status = "成功"
    note = ""
    try:
        # 1) 核价 + 加工副本 + 生成天图核验清单
        if not args.skip_engine:
            extra = ["--no-write-dialog", "--input", bill_arg]
            code, secs = run_stage("核价引擎", HERE / "run_yongda_check.py", extra)
            stage_reports.append({"name": "核价引擎", "secs": secs, "ok": code == 0})
            if code != 0:
                status = "失败"
                note = f"核价引擎退出码 {code}"
                raise SystemExit(code)
        checklist = latest(out_dir, "*_天图核验清单_*.csv")

        # 2) 天图核验
        if not args.skip_tiantu:
            extra = []
            if checklist:
                extra += ["--csv", str(checklist)]
            if args.tiantu_limit:
                extra += ["--limit", str(args.tiantu_limit)]
            code, secs = run_stage("天图核验", HERE / "check_tiantu.py", extra)
            stage_reports.append({"name": "天图核验", "secs": secs, "ok": code == 0})
            if code != 0:
                status = "失败"
                note = f"天图核验退出码 {code}"
                raise SystemExit(code)
        tiantu_result = latest(out_dir, "天图核验结果_*.csv")

        # 3) 合并问题，弹窗写入统计表
        if not args.skip_write:
            extra = ["--dry-run"] if args.dry_run else []
            if tiantu_result:
                extra += ["--tiantu-result", str(tiantu_result)]
            if checklist:
                extra += ["--checklist", str(checklist)]
            code, secs = run_stage("写入问题统计表", HERE / "write_problems_to_stats.py", extra)
            stage_reports.append({"name": "写入统计表", "secs": secs, "ok": code == 0})
            if code != 0:
                status = "失败"
                note = f"写入问题统计表退出码 {code}"
                raise SystemExit(code)

        # 汇总本批产物
        print("\n" + "=" * 60)
        print("[主流程] 全部完成，本次产物：")
        names = [
            ("加工副本(含四列)", "*_加工副本_*.xlsx"),
            ("核价结果", "*_核价结果_*.xlsx"),
            ("天图核验清单", "*_天图核验清单_*.csv"),
            ("天图核验结果", "天图核验结果_*.csv"),
        ]
        for label, pattern in names:
            f = latest(out_dir, pattern)
            print(f"  {label}: {f if f else '未生成'}")
    except SystemExit:
        raise
    except Exception as exc:  # noqa: BLE001
        status = "失败"
        note = f"主流程异常: {exc}"
        raise
    finally:
        notify_result(status, start_ts, stage_reports, note)


if __name__ == "__main__":
    main()
