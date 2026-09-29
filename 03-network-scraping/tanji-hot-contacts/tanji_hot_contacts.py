#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""探迹 CRM Hot 联系方式补全。

浏览器依赖在真正运行抓取时才导入，因此 --validate-only 和离线 OCR 测试
可在未安装 Selenium 的环境中执行。
"""

from __future__ import annotations

import argparse
import base64
import csv
import getpass
import io
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from xml.etree import ElementTree
from zipfile import ZipFile
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageEnhance, ImageFilter, ImageOps


CUSTOMER_HEADERS = {"客户名", "客户名称"}
CONTACT_HEADERS = {"Hot联系方式", "HOT联系方式", "hot联系方式"}
MOBILE_RE = re.compile(r"1[3-9]\d{9}")
PHONE_RE = re.compile(r"(?<!\d)(?:\d{3,4}-?)?\d{7,11}(?!\d)")


@dataclass(frozen=True)
class InputRecord:
    source_row: int
    customer_name: str
    masked_contact: str


@dataclass
class ContactResult:
    customer_name: str
    hot_contact: str
    status: str
    source_row: int
    masked_contact: str
    contact_name: str = ""
    confidence: str = ""
    note: str = ""
    captured_at: str = ""

    def __post_init__(self) -> None:
        if not self.captured_at:
            self.captured_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")


@dataclass(frozen=True)
class Settings:
    browser: str
    chrome_binary_path: str
    login_url: str
    customer_list_url: str
    page_timeout_seconds: int
    search_timeout_seconds: int
    reveal_timeout_seconds: int
    retry_count: int
    checkpoint_every: int
    tesseract_path: str


def norm_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def load_settings(path: Path) -> Settings:
    raw = json.loads(path.read_text(encoding="utf-8"))
    return Settings(**raw)


def read_credentials_from_docx(path: Path) -> tuple[str, str]:
    """从用户指定的 SOP DOCX 内存读取账号密码，不落盘、不记录日志。"""
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    text = "\n".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
    username = re.search(r"账号[:：]\s*([^\s]+)", text)
    password = re.search(r"密码[:：]\s*([^\s]+)", text)
    if not username or not password:
        raise ValueError("指定 DOCX 中未找到‘账号：’或‘密码：’")
    return username.group(1), password.group(1)


def find_input_columns(ws) -> tuple[int, int, int]:
    for row_no in range(1, min(ws.max_row, 10) + 1):
        customer_col = contact_col = None
        for col_no in range(1, ws.max_column + 1):
            value = norm_text(ws.cell(row_no, col_no).value)
            if value in CUSTOMER_HEADERS:
                customer_col = col_no
            if value in CONTACT_HEADERS:
                contact_col = col_no
        if customer_col and contact_col:
            return row_no, customer_col, contact_col
    raise ValueError("前 10 行未同时找到‘客户名/客户名称’和‘Hot联系方式’表头")


def read_input_records(path: Path, start_row: int = 0, limit: int | None = None) -> list[InputRecord]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header_row, customer_col, contact_col = find_input_columns(ws)
        first_data_row = max(header_row + 1, start_row or header_row + 1)
        records: list[InputRecord] = []
        for row_no, values in enumerate(ws.iter_rows(min_row=first_data_row, values_only=True), first_data_row):
            customer = str(values[customer_col - 1] or "").strip() if customer_col <= len(values) else ""
            masked = str(values[contact_col - 1] or "").strip() if contact_col <= len(values) else ""
            if not customer:
                continue
            records.append(InputRecord(row_no, customer, masked))
            if limit is not None and len(records) >= limit:
                break
        return records
    finally:
        wb.close()


def mask_pattern(masked: str) -> re.Pattern[str] | None:
    compact = re.sub(r"[\s()（）-]", "", masked or "")
    if not re.search(r"[*＊xX•·]", compact) or not re.search(r"\d", compact):
        return None
    parts = []
    for char in compact:
        if char.isdigit():
            parts.append(re.escape(char))
        elif char in "*＊xX•·":
            parts.append(r"\d")
    return re.compile("^" + "".join(parts) + "$") if parts else None


def split_masked_contacts(masked: str) -> list[str]:
    """拆分源表中可能用逗号、分号或竖线拼接的多个掩码。"""
    values = [part.strip() for part in re.split(r"[,，;；|、/\\\n]+", masked or "") if part.strip()]
    return values or ([masked.strip()] if masked and masked.strip() else [])


def normalize_ocr_text(text: str) -> list[str]:
    text = text.translate(str.maketrans({"O": "0", "o": "0", "I": "1", "l": "1", "|": "1"}))
    candidates = []
    for raw in re.findall(r"[\d()（）+\-\s]{7,}", text):
        digits = re.sub(r"\D", "", raw)
        if 7 <= len(digits) <= 15:
            candidates.append(digits)
        if len(digits) > 11:
            candidates.extend(MOBILE_RE.findall(digits))
    return candidates


class OCRReader:
    def __init__(self, tesseract_path: str):
        self.tesseract = Path(tesseract_path)
        if not self.tesseract.exists():
            resolved = shutil.which(tesseract_path) or shutil.which("tesseract")
            if not resolved:
                raise FileNotFoundError(f"未找到 Tesseract: {tesseract_path}")
            self.tesseract = Path(resolved)

    @staticmethod
    def decode_data_url(data_url: str) -> Image.Image:
        if "," not in data_url:
            raise ValueError("图片不是 data URL")
        payload = data_url.split(",", 1)[1]
        image = Image.open(io.BytesIO(base64.b64decode(payload)))
        if image.mode == "RGBA":
            background = Image.new("RGBA", image.size, "white")
            image = Image.alpha_composite(background, image).convert("RGB")
        return image.convert("RGB")

    @staticmethod
    def variants(image: Image.Image) -> Iterable[tuple[str, Image.Image]]:
        scale = max(6, min(12, 480 // max(image.width, 1)))
        enlarged = image.resize((image.width * scale, image.height * scale), Image.Resampling.LANCZOS)
        gray = ImageOps.autocontrast(enlarged.convert("L"))
        yield "gray", gray
        yield "sharp", ImageEnhance.Contrast(gray.filter(ImageFilter.SHARPEN)).enhance(2.0)
        for threshold in (110, 140, 170, 200):
            yield f"threshold_{threshold}", gray.point(lambda p, t=threshold: 255 if p > t else 0)

    def _run(self, image: Image.Image, psm: int) -> str:
        temp_path = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
                temp_path = Path(handle.name)
            image.save(temp_path)
            command = [
                str(self.tesseract), str(temp_path), "stdout", "-l", "eng",
                "--oem", "3", "--psm", str(psm),
                "-c", "tessedit_char_whitelist=0123456789-()",
            ]
            proc = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=15)
            return proc.stdout.strip() if proc.returncode == 0 else ""
        finally:
            if temp_path:
                temp_path.unlink(missing_ok=True)

    def recognize(self, data_url: str, masked: str = "") -> tuple[str, str, str, Image.Image]:
        image = self.decode_data_url(data_url)
        observations: list[str] = []
        for _, variant in self.variants(image):
            for psm in (7, 13):
                observations.extend(normalize_ocr_text(self._run(variant, psm)))

        if not observations:
            return "", "低", "未识别到有效数字", image

        counts = Counter(observations)
        pattern = mask_pattern(masked)

        def score(item: tuple[str, int]) -> tuple[int, int, int]:
            candidate, votes = item
            mask_bonus = 20 if pattern and pattern.fullmatch(candidate) else 0
            mobile_bonus = 5 if MOBILE_RE.fullmatch(candidate) else 0
            return mask_bonus + mobile_bonus + votes * 2, votes, len(candidate)

        best, votes = max(counts.items(), key=score)
        matches_mask = bool(pattern and pattern.fullmatch(best))
        valid_phone = bool(MOBILE_RE.fullmatch(best) or PHONE_RE.fullmatch(best))
        if matches_mask and votes >= 1:
            return best, "高" if votes >= 2 else "中", f"{votes} 个 OCR 结果支持且与掩码一致", image
        if valid_phone and votes >= 3:
            note = f"{votes} 个 OCR 结果支持"
            if pattern:
                note += "，但与输入掩码不一致"
            return best, "中", note, image
        return best if valid_phone else "", "低", f"OCR 证据不足；候选={best}，票数={votes}", image


def xpath_literal(text: str) -> str:
    if "'" not in text:
        return f"'{text}'"
    if '"' not in text:
        return f'"{text}"'
    pieces = text.split("'")
    return "concat(" + ", \"'\", ".join(f"'{p}'" for p in pieces) + ")"


class TanjiBrowser:
    def __init__(self, settings: Settings, evidence_dir: Path, headless: bool = False, browser: str | None = None):
        try:
            from selenium import webdriver
        except ImportError as exc:
            raise SystemExit("缺少 selenium，请先安装 requirements.txt") from exc

        browser_name = (browser or settings.browser or "chrome").lower()
        if browser_name not in {"chrome", "edge"}:
            raise ValueError(f"不支持的浏览器: {browser_name}，可选 chrome 或 edge")
        if browser_name == "chrome":
            from selenium.webdriver.chrome.options import Options
        else:
            from selenium.webdriver.edge.options import Options

        options = Options()
        if headless:
            options.add_argument("--headless=new")
        else:
            options.add_argument("--start-maximized")
        options.add_argument("--disable-notifications")
        options.add_argument("--disable-blink-features=AutomationControlled")
        options.add_experimental_option("excludeSwitches", ["enable-logging", "enable-automation"])
        options.add_experimental_option("useAutomationExtension", False)
        if browser_name == "chrome":
            if settings.chrome_binary_path and Path(settings.chrome_binary_path).exists():
                options.binary_location = settings.chrome_binary_path
            self.driver = webdriver.Chrome(options=options)
        else:
            self.driver = webdriver.Edge(options=options)
        self.driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
        self.driver.set_page_load_timeout(settings.page_timeout_seconds)
        self.settings = settings
        self.evidence_dir = evidence_dir

    def close(self) -> None:
        try:
            self.driver.quit()
        except Exception:
            pass

    def save_page_evidence(self, tag: str) -> None:
        safe = re.sub(r"[^\w.-]+", "_", tag)[:80]
        try:
            self.driver.save_screenshot(str(self.evidence_dir / f"{safe}.png"))
            (self.evidence_dir / f"{safe}.html").write_text(self.driver.page_source, encoding="utf-8")
        except Exception as exc:
            logging.warning("保存页面证据失败: %s", exc)

    def login(self, username: str, password: str, manual: bool) -> None:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.common.keys import Keys
        from selenium.webdriver.support.ui import WebDriverWait

        self.driver.get(self.settings.login_url)
        wait = WebDriverWait(self.driver, self.settings.page_timeout_seconds)
        if manual:
            input("请在浏览器完成登录，进入系统后回到终端按 Enter：")
        else:
            notice_buttons = self.driver.find_elements(
                By.XPATH,
                "//button[contains(@class,'ant-btn-primary')][.//span[contains(normalize-space(.),'我知道了')]]",
            )
            notice_buttons += self.driver.find_elements(By.XPATH, "//*[self::button or @role='button'][contains(normalize-space(.),'我知道了')]")
            for elem in notice_buttons:
                if elem.is_displayed():
                    try:
                        elem.click()
                        break
                    except Exception:
                        pass
            for elem in self.driver.find_elements(By.XPATH, "//*[self::button or @role='button' or self::a][contains(normalize-space(.),'账号登录')]"):
                if elem.is_displayed():
                    try:
                        elem.click()
                        break
                    except Exception:
                        pass
            try:
                user_box = wait.until(lambda d: next((x for x in d.find_elements(By.CSS_SELECTOR, "input[type='text'],input[type='tel']") if x.is_displayed()), None))
                pass_box = wait.until(lambda d: next((x for x in d.find_elements(By.CSS_SELECTOR, "input[type='password']") if x.is_displayed()), None))
            except Exception:
                self.save_page_evidence("login_form_unavailable")
                input("自动登录控件未出现（可能是登录弹窗、验证码或会话提示）。请在浏览器手动登录并进入客户列表，完成后回到终端按 Enter：")
                self.driver.get(self.settings.customer_list_url)
                wait.until(lambda d: d.find_elements(By.CSS_SELECTOR, "table") or d.find_elements(By.XPATH, "//*[contains(normalize-space(.),'客户管理')]"))
                return
            user_box.clear()
            user_box.send_keys(username)
            pass_box.clear()
            pass_box.send_keys(password)
            login_buttons = self.driver.find_elements(
                By.XPATH,
                "//button[contains(normalize-space(.),'登录') or .//span[contains(normalize-space(.),'登录')]]",
            )
            clickable = next((item for item in login_buttons if item.is_displayed() and item.is_enabled()), None)
            if clickable:
                clickable.click()
            else:
                pass_box.send_keys(Keys.ENTER)

        wait.until(lambda d: "/login" not in d.current_url or d.find_elements(By.XPATH, "//*[contains(normalize-space(.),'客户管理')]"))
        body_text = self.driver.find_element(By.TAG_NAME, "body").text
        if "请先选择登录企业" in body_text:
            input("账号密码已自动提交。请在浏览器点击‘我知道了’，选择登录企业并进入系统；完成后告知操作端继续：")
        self.driver.get(self.settings.customer_list_url)
        try:
            wait.until(lambda d: d.find_elements(By.CSS_SELECTOR, "table") or d.find_elements(By.XPATH, "//*[contains(normalize-space(.),'客户管理')]"))
        except Exception:
            body_text = self.driver.find_element(By.TAG_NAME, "body").text
            reason = "需要选择登录企业" if "请先选择登录企业" in body_text else "尚未进入客户列表"
            self.save_page_evidence("login_takeover")
            input(f"账号密码已自动提交，但{reason}。请在浏览器完成操作并进入客户列表；完成后告知操作端继续：")
            self.driver.get(self.settings.customer_list_url)
            wait.until(lambda d: d.find_elements(By.CSS_SELECTOR, "table") or d.find_elements(By.XPATH, "//*[contains(normalize-space(.),'客户管理')]"))

    def _search_box(self):
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait

        selectors = [
            "input[placeholder*='搜索客户名称']",
            "input[placeholder*='客户名称']",
            "input[placeholder*='搜索']",
        ]
        wait = WebDriverWait(self.driver, self.settings.search_timeout_seconds)
        return wait.until(lambda d: next((d.find_element(By.CSS_SELECTOR, s) for s in selectors if d.find_elements(By.CSS_SELECTOR, s) and d.find_element(By.CSS_SELECTOR, s).is_displayed()), None))

    def search(self, customer_name: str) -> None:
        from selenium.webdriver.common.keys import Keys
        from selenium.webdriver.support.ui import WebDriverWait

        box = self._search_box()
        old_rows = "\n".join(self._visible_row_texts())
        box.click()
        box.send_keys(Keys.CONTROL, "a")
        box.send_keys(customer_name)
        box.send_keys(Keys.ENTER)

        def refreshed(_driver):
            current = "\n".join(self._visible_row_texts())
            body = self.driver.find_element("tag name", "body").text
            return current and (current != old_rows or customer_name in current) and "加载中" not in body

        WebDriverWait(self.driver, self.settings.search_timeout_seconds).until(refreshed)

    def _visible_row_texts(self) -> list[str]:
        from selenium.webdriver.common.by import By

        rows = self.driver.find_elements(By.CSS_SELECTOR, "tr.ant4-table-row, tbody tr")
        return [re.sub(r"\s+", " ", row.text).strip() for row in rows if row.is_displayed()]

    def _exact_customer_row(self, customer_name: str):
        from selenium.webdriver.common.by import By

        expected = norm_text(customer_name)
        for row in self.driver.find_elements(By.CSS_SELECTOR, "tr.ant4-table-row, tbody tr"):
            if not row.is_displayed():
                continue
            cells = row.find_elements(By.CSS_SELECTOR, "td")
            if any(norm_text(cell.text) == expected for cell in cells):
                return row
            links = row.find_elements(By.CSS_SELECTOR, "a")
            if any(norm_text(link.text) == expected for link in links):
                return row
        return None

    def _hot_column_index(self) -> int | None:
        from selenium.webdriver.common.by import By

        headers = self.driver.find_elements(By.CSS_SELECTOR, "thead th")
        for index, header in enumerate(headers):
            if "hot联系方式" in norm_text(header.text).lower():
                return index
        return None

    def _hot_cell(self, row):
        from selenium.webdriver.common.by import By

        cells = row.find_elements(By.CSS_SELECTOR, "td")
        index = self._hot_column_index()
        if index is not None and index < len(cells):
            return cells[index]
        for cell in cells:
            if cell.find_elements(By.XPATH, ".//i[.//img[contains(@src,'.svg')]]"):
                return cell
        return None

    @staticmethod
    def _contact_blocks(cell):
        from selenium.webdriver.common.by import By

        blocks = cell.find_elements(By.XPATH, "./div[.//img[starts-with(@src,'data:image')]]")
        return blocks or cell.find_elements(By.XPATH, ".//*[self::div or self::span][.//img[starts-with(@src,'data:image')]]")

    @staticmethod
    def _eye(block):
        from selenium.webdriver.common.by import By

        preferred = block.find_elements(By.XPATH, ".//i[.//img[contains(@src,'ca2d2e375c4ad3baaf7c5d154e34c5a9')]]")
        if preferred:
            return preferred[0]
        fallback = block.find_elements(By.XPATH, ".//i[.//img[contains(@src,'.svg') and not(contains(@src,'icon-phone'))]]")
        return fallback[0] if fallback else None

    def extract(self, record: InputRecord, ocr: OCRReader, save_images: str) -> list[ContactResult]:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait

        self.search(record.customer_name)
        row = self._exact_customer_row(record.customer_name)
        if row is None:
            return [ContactResult(record.customer_name, "", "未找到客户", record.source_row, record.masked_contact)]
        cell = self._hot_cell(row)
        if cell is None:
            return [ContactResult(record.customer_name, "", "页面错误", record.source_row, record.masked_contact, note="未识别到 Hot 联系方式列")]
        blocks = self._contact_blocks(cell)
        if not blocks:
            return [ContactResult(record.customer_name, "", "未找到查看按钮", record.source_row, record.masked_contact, note="Hot 联系方式单元格中没有号码图片")]

        results: list[ContactResult] = []
        for block_index in range(len(blocks)):
            row = self._exact_customer_row(record.customer_name)
            cell = self._hot_cell(row)
            blocks = self._contact_blocks(cell)
            if block_index >= len(blocks):
                results.append(ContactResult(record.customer_name, "", "页面错误", record.source_row, record.masked_contact, note=f"第 {block_index + 1} 个联系方式刷新后消失"))
                continue
            block = blocks[block_index]
            label_text = block.text.strip()
            contact_name = re.split(r"[:：]", label_text, maxsplit=1)[0].strip()
            number_images = block.find_elements(By.XPATH, ".//img[starts-with(@src,'data:image')]")
            before_src = number_images[0].get_attribute("src") if number_images else ""
            eye = self._eye(block)
            if eye is None:
                results.append(ContactResult(record.customer_name, "", "未找到查看按钮", record.source_row, record.masked_contact, contact_name=contact_name))
                continue
            self.driver.execute_script("arguments[0].scrollIntoView({block:'center'}); arguments[0].click();", eye)

            def changed(_driver):
                fresh_row = self._exact_customer_row(record.customer_name)
                fresh_cell = self._hot_cell(fresh_row) if fresh_row else None
                fresh_blocks = self._contact_blocks(fresh_cell) if fresh_cell else []
                if block_index >= len(fresh_blocks):
                    return False
                images = fresh_blocks[block_index].find_elements(By.XPATH, ".//img[starts-with(@src,'data:image')]")
                return images[0].get_attribute("src") if images and images[0].get_attribute("src") != before_src else False

            try:
                data_url = WebDriverWait(self.driver, self.settings.reveal_timeout_seconds).until(changed)
                masks = split_masked_contacts(record.masked_contact)
                masked_for_block = masks[block_index] if block_index < len(masks) else ""
                phone, confidence, note, original_image = ocr.recognize(data_url, masked_for_block)
                status = "成功" if phone and confidence in {"高", "中"} and "不一致" not in note else ("需复核" if phone else "OCR失败")
                if save_images == "all" or status != "成功":
                    safe = re.sub(r"[^\w.-]+", "_", record.customer_name)[:50]
                    original_image.save(self.evidence_dir / f"row_{record.source_row}_{safe}_{block_index + 1}.png")
                results.append(ContactResult(record.customer_name, phone, status, record.source_row, record.masked_contact, contact_name, confidence, note))
            except Exception as exc:
                results.append(ContactResult(record.customer_name, "", "页面错误", record.source_row, record.masked_contact, contact_name=contact_name, note=f"点击后图片未变化: {exc}"))

        return results


def load_checkpoint(path: Path) -> dict[int, list[ContactResult]]:
    done: dict[int, list[ContactResult]] = {}
    if not path.exists():
        return done
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            try:
                event = json.loads(line)
                done[int(event["source_row"])] = [ContactResult(**item) for item in event["results"]]
            except (json.JSONDecodeError, KeyError, TypeError, ValueError):
                logging.warning("忽略损坏的断点行")
    return done


def append_checkpoint(path: Path, source_row: int, results: list[ContactResult]) -> None:
    event = {"source_row": source_row, "results": [asdict(item) for item in results]}
    with path.open("a", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(event, ensure_ascii=False) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def write_results(path: Path, result_map: dict[int, list[ContactResult]]) -> Path | None:
    wb = Workbook()
    ws = wb.active
    ws.title = "抓取结果"
    headers = ["客户名称", "Hot联系方式", "状态", "联系人", "OCR置信度", "原始行号", "原始掩码", "说明", "抓取时间"]
    ws.append(headers)
    for source_row in sorted(result_map):
        for item in result_map[source_row]:
            ws.append([item.customer_name, item.hot_contact, item.status, item.contact_name, item.confidence, item.source_row, item.masked_contact, item.note, item.captured_at])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = [38, 20, 14, 16, 12, 12, 20, 46, 20]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
        row[5].alignment = Alignment(horizontal="center", vertical="center")
    ws.sheet_view.showGridLines = False

    summary = wb.create_sheet("运行汇总")
    flat = [item for values in result_map.values() for item in values]
    counts = Counter(item.status for item in flat)
    summary.append(["指标", "数量"])
    summary.append(["已处理输入行", len(result_map)])
    summary.append(["输出联系方式行", len(flat)])
    for status in ("成功", "需复核", "未找到客户", "未找到查看按钮", "OCR失败", "页面错误"):
        summary.append([status, counts.get(status, 0)])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 14
    summary.sheet_view.showGridLines = False

    temp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(temp)
    try:
        os.replace(temp, path)
        return path
    except PermissionError:
        # 用户查看结果时 Excel 会锁定主文件。抓取不能因此中止；将最新结果
        # 写入固定的“运行中”副本，下一次刷新继续替换该副本。
        fallback = path.with_name(path.stem + "_运行中.xlsx")
        try:
            os.replace(temp, fallback)
            logging.warning("主结果表被占用，最新结果已改写到：%s", fallback)
            return fallback
        except PermissionError:
            logging.warning("主结果表和运行中副本均被占用，本次跳过 Excel 刷新；JSONL 断点不受影响")
            temp.unlink(missing_ok=True)
            return None


def setup_logging(run_dir: Path) -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=[logging.FileHandler(run_dir / "run.log", encoding="utf-8"), logging.StreamHandler(sys.stdout)],
        force=True,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="根据 Excel 补全探迹 CRM Hot 联系方式")
    parser.add_argument("--input", type=Path)
    parser.add_argument("--config", type=Path, default=Path(__file__).with_name("config.json"))
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--start-row", type=int, default=0)
    parser.add_argument("--limit", type=int)
    parser.add_argument("--manual-login", action="store_true")
    parser.add_argument("--credential-docx", type=Path, help="从指定 SOP DOCX 内存读取登录凭据")
    parser.add_argument("--headless", action="store_true")
    parser.add_argument("--browser", choices=("chrome", "edge"), help="浏览器，默认读取 config.json")
    parser.add_argument("--save-images", choices=("failures", "all"), default="failures")
    parser.add_argument("--no-resume", action="store_true")
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--ocr-image", type=Path, help="离线测试单张号码图片")
    parser.add_argument("--masked", default="", help="配合 --ocr-image 使用的掩码")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    settings = load_settings(args.config.resolve())

    if args.ocr_image:
        ocr = OCRReader(settings.tesseract_path)
        payload = "data:image/png;base64," + base64.b64encode(args.ocr_image.read_bytes()).decode("ascii")
        phone, confidence, note, _ = ocr.recognize(payload, args.masked)
        print(json.dumps({"号码": phone, "置信度": confidence, "说明": note}, ensure_ascii=False))
        return 0 if phone else 2

    if args.input is None:
        raise SystemExit("常规抓取必须提供 --input；离线 OCR 可只提供 --ocr-image")

    input_path = args.input.resolve()
    if not input_path.exists():
        raise FileNotFoundError(input_path)
    records = read_input_records(input_path, args.start_row, args.limit)
    print(f"输入检查通过：{input_path.name}，待处理 {len(records)} 行")
    if args.validate_only:
        return 0

    if args.output_dir:
        run_dir = args.output_dir.resolve()
    else:
        run_dir = Path(__file__).with_name("output") / f"run_{datetime.now():%Y%m%d_%H%M%S}"
    run_dir.mkdir(parents=True, exist_ok=True)
    evidence_dir = run_dir / "evidence"
    evidence_dir.mkdir(exist_ok=True)
    setup_logging(run_dir)

    checkpoint_path = run_dir / "checkpoint.jsonl"
    if args.no_resume and checkpoint_path.exists():
        checkpoint_path = run_dir / f"checkpoint_{datetime.now():%Y%m%d_%H%M%S}.jsonl"
    result_map = load_checkpoint(checkpoint_path)
    output_path = run_dir / "探迹Hot联系方式结果.xlsx"
    logging.info("输入=%s；总行数=%d；已完成=%d", input_path.name, len(records), len(result_map))

    username = os.environ.get("TANZHI_USERNAME", "").strip()
    password = os.environ.get("TANZHI_PASSWORD", "")
    if not args.manual_login:
        if (not username or not password) and args.credential_docx:
            username, password = read_credentials_from_docx(args.credential_docx.resolve())
        username = username or input("探迹 CRM 账号：").strip()
        password = password or getpass.getpass("探迹 CRM 密码：")
        if not username or not password:
            raise ValueError("账号和密码不能为空；也可使用 --manual-login")

    ocr = OCRReader(settings.tesseract_path)
    browser = TanjiBrowser(settings, evidence_dir, args.headless, args.browser)
    try:
        browser.login(username, password, args.manual_login)
        for position, record in enumerate(records, 1):
            if record.source_row in result_map:
                continue
            logging.info("[%d/%d] Excel 行 %d：%s", position, len(records), record.source_row, record.customer_name)
            last_error = None
            results = None
            for attempt in range(settings.retry_count + 1):
                try:
                    results = browser.extract(record, ocr, args.save_images)
                    break
                except Exception as exc:
                    last_error = exc
                    logging.warning("第 %d 次尝试失败：%s", attempt + 1, exc)
                    time.sleep(min(2 ** attempt, 5))
            if results is None:
                results = [ContactResult(record.customer_name, "", "页面错误", record.source_row, record.masked_contact, note=str(last_error))]
                browser.save_page_evidence(f"row_{record.source_row}_page_error")
            elif any(item.status != "成功" for item in results):
                browser.save_page_evidence(f"row_{record.source_row}_{results[0].status}")
            result_map[record.source_row] = results
            append_checkpoint(checkpoint_path, record.source_row, results)
            if len(result_map) % settings.checkpoint_every == 0:
                written_path = write_results(output_path, result_map)
                if written_path:
                    logging.info("已刷新 Excel 断点：%s", written_path)
    except KeyboardInterrupt:
        logging.warning("收到中断，正在保存当前结果")
    finally:
        browser.close()
        write_results(output_path, result_map)

    counts = Counter(item.status for values in result_map.values() for item in values)
    logging.info("完成：%s；状态=%s", output_path, dict(counts))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
